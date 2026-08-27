import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def add_label_to_placeholder(slide, right_ph, label_text):
    # Add a small text box above the placeholder
    left = right_ph.left
    top = right_ph.top - Inches(0.4)
    width = right_ph.width
    height = Inches(0.4)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = label_text
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(16, 185, 129) # Emerald Green to match theme
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

def create_milestone2_slides():
    prs = Presentation('ShopSense_Milestone1_Final.pptx')
    
    # Delete slides 10 and 11 (Future Enhancements and Thank You)
    # The length should be 12.
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    for slide in slides[10:]:
        prs.part.drop_rel(slide.rId)
        xml_slides.remove(slide)
        
    layout_1 = prs.slide_layouts[1] # Title and Content
    layout_3 = prs.slide_layouts[3] # Two Content
    
    # PAGE 11
    slide11 = prs.slides.add_slide(layout_1)
    slide11.shapes.title.text = "Milestone 2: Inventory Intelligence & Customer Analytics"
    tf11 = slide11.placeholders[1].text_frame
    tf11.text = "Milestone 2 extends the marketplace with real customer transactions and data-driven analytics."
    bullets_11 = [
        "Extended ShopSense with a customer-facing marketplace and real order flow.",
        "Connected customer purchases with transaction and inventory data.",
        "Added customer segmentation and sales-based product recommendations.",
        "Added AI features for demand forecasting and review sentiment analysis."
    ]
    for b in bullets_11:
        p = tf11.add_paragraph()
        p.text = b
        p.level = 0
        
    # PAGE 12
    slide12 = prs.slides.add_slide(layout_3)
    slide12.shapes.title.text = "Customer Portal & Transactions"
    tf12 = slide12.placeholders[1].text_frame
    tf12.text = "Customer enters basic details to start shopping."
    bullets_12 = [
        "Browses the marketplace and places an order.",
        "The system stores the transaction automatically.",
        "Product stock is instantly reduced in the PostgreSQL database."
    ]
    for b in bullets_12:
        p = tf12.add_paragraph()
        p.text = b
        p.level = 0
    add_label_to_placeholder(slide12, slide12.placeholders[2], "Insert Customer Portal / Marketplace Screenshot Here ↓")

    # PAGE 13
    slide13 = prs.slides.add_slide(layout_3)
    slide13.shapes.title.text = "Customer Analytics & Segmentation"
    tf13 = slide13.placeholders[1].text_frame
    tf13.text = "Customer spending is calculated from transaction history and customers are grouped into:"
    bullets_13 = [
        "High Value: $500+",
        "Regular: $100–$499",
        "Low Value: below $100",
        "Analytics are calculated directly from actual transaction records."
    ]
    for b in bullets_13:
        p = tf13.add_paragraph()
        p.text = b
        p.level = 0 if "Value:" not in b and "Regular:" not in b else 1
    add_label_to_placeholder(slide13, slide13.placeholders[2], "Insert Customer Analytics Screenshot Here ↓")

    # PAGE 14
    slide14 = prs.slides.add_slide(layout_3)
    slide14.shapes.title.text = "Gemini Review Sentiment Analysis"
    tf14 = slide14.placeholders[1].text_frame
    tf14.text = "Vendor selects a product."
    bullets_14 = [
        "Enters rating and customer review text.",
        "Gemini AI analyzes the review.",
        "The system returns a sentiment score and identifies positive feedback and pros.",
        "This helps the vendor quickly understand customer opinions."
    ]
    for b in bullets_14:
        p = tf14.add_paragraph()
        p.text = b
        p.level = 0
    add_label_to_placeholder(slide14, slide14.placeholders[2], "Insert Gemini Sentiment Analysis Screenshot Here ↓")

    # PAGE 15
    slide15 = prs.slides.add_slide(layout_3)
    slide15.shapes.title.text = "Rule-Based Recommendations & Historical Validation"
    tf15 = slide15.placeholders[1].text_frame
    tf15.text = "Product sales are calculated from recorded transaction data."
    bullets_15 = [
        "Products are ranked based on the quantity sold.",
        "The highest-selling products are shown as recommendations.",
        "Historical/test transaction data was used to verify that the analytics produce the expected results."
    ]
    for b in bullets_15:
        p = tf15.add_paragraph()
        p.text = b
        p.level = 0
    add_label_to_placeholder(slide15, slide15.placeholders[2], "Insert Recommendation / Historical Validation Screenshot Here ↓")

    # PAGE 16
    slide16 = prs.slides.add_slide(layout_3)
    slide16.shapes.title.text = "ARIMA Demand Forecasting"
    tf16 = slide16.placeholders[1].text_frame
    tf16.text = "Historical sales data is used to identify demand patterns."
    bullets_16 = [
        "The vendor selects a product.",
        "ARIMA generates a 7-day demand forecast.",
        "The forecast is compared with current inventory.",
        "A restock alert is shown when expected demand is higher than available stock."
    ]
    for b in bullets_16:
        p = tf16.add_paragraph()
        p.text = b
        p.level = 0
    add_label_to_placeholder(slide16, slide16.placeholders[2], "Insert ARIMA Forecast Screenshot Here ↓")

    # PAGE 17
    slide17 = prs.slides.add_slide(layout_3)
    slide17.shapes.title.text = "Milestone 2 Outcome"
    tf17 = slide17.placeholders[1].text_frame
    tf17.text = "Customer purchases now generate real transaction data, which feeds the analytics layer. This data is used for customer segmentation, product recommendations, demand forecasting, and review sentiment analysis."
    bullets_17 = [
        "Real customer ordering and transaction tracking",
        "Data-driven customer segmentation and recommendations",
        "7-day inventory demand forecasting with ARIMA",
        "Gemini-powered review sentiment analysis"
    ]
    for b in bullets_17:
        p = tf17.add_paragraph()
        p.text = b
        p.level = 0
    add_label_to_placeholder(slide17, slide17.placeholders[2], "Insert Final Analytics Dashboard Screenshot Here ↓")
    
    prs.save('ShopSense_Final_Presentation_Revised.pptx')
    print("Saved ShopSense_Final_Presentation_Revised.pptx")

if __name__ == '__main__':
    create_milestone2_slides()
